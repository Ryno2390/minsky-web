"""BEA fixed assets by INDUSTRY and asset type, from the agency's own detail files.

FRED carries BEA's fixed assets only in economy-wide aggregates, which is why
hollowing.py had to answer a question about manufacturing using economy-wide numbers
and hedged its verdict accordingly. BEA publishes the underlying detail as spreadsheets
that need no API key:

    https://apps.bea.gov/national/FA2004/Details/xls/

One workbook per measure, one sheet per industry, one row per asset type, one column per
year from 1925. 76 industries at roughly NAICS three-digit level.

    detailnonres_stk1.xlsx   current-cost net stock
    detailnonres_stk2.xlsx   the same stock in chained 2017 dollars, so that stk1/stk2
                             is an asset-specific price index consistent with both
    detailnonres_dep1.xlsx   current-cost depreciation
    DetailNonres_rate.xlsx   BEA's own implied depreciation rates

Each sheet carries three aggregate rows -- TOTAL EQUIPMENT, TOTAL STRUCTURES, TOTAL
INTELLECTUAL PROPERTY PRODUCTS -- and those are what this extracts. The asset-level
detail beneath them is left in the workbook.

WHY IT CACHES TWICE
-------------------
The workbooks are 10 MB each and parsing 76 sheets takes long enough to be annoying in a
loop, so the extraction is written to a small CSV beside them and read from there
afterwards. Both live in the gitignored cache directory. Deleting the CSV forces a
re-parse; deleting the xlsx forces a re-download.
"""
import csv
import io
import subprocess
from pathlib import Path

CACHE = Path(__file__).parent / "cache"
CACHE.mkdir(exist_ok=True)

BASE = "https://apps.bea.gov/national/FA2004/Details/xls"
FILES = {"stock": "detailnonres_stk1.xlsx",      # current-cost net stock
         "real": "detailnonres_stk2.xlsx",       # same, in chained 2017 dollars
         "deprec": "detailnonres_dep1.xlsx",     # current-cost depreciation
         "rate": "DetailNonres_rate.xlsx"}       # BEA's own implied rates

#: the three aggregate rows present on every industry sheet
TOTALS = {"EQUIPMENT": "equipment", "STRUCTURES": "structures", "IPP": "ip"}

#: apps.bea.gov refuses a bare urllib request; curl with a browser agent is served.
UA = ("Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/120 Safari/537.36")


def _download(which):
    path = CACHE / FILES[which]
    if path.exists() and path.stat().st_size > 1_000_000:
        return path
    url = f"{BASE}/{FILES[which]}"
    out = subprocess.run(["curl", "-sL", "--max-time", "300", "-A", UA, "-o", str(path),
                          url], capture_output=True, timeout=400)
    if out.returncode != 0 or not path.exists():
        raise RuntimeError(f"could not fetch {url}")
    head = path.read_bytes()[:4]
    if head[:2] != b"PK":                      # an xlsx is a zip; HTML means refused
        path.unlink()
        raise RuntimeError(f"{url} returned HTML rather than a workbook")
    return path


def _year(cell):
    """1925 from 1925 or '1925'; None from anything else."""
    try:
        y = int(str(cell).strip())
    except (TypeError, ValueError):
        return None
    return y if 1900 <= y <= 2100 else None


def _extract(which):
    """[(industry_code, industry_name, asset, year, value)] from the workbook."""
    import openpyxl
    wb = openpyxl.load_workbook(_download(which), read_only=True, data_only=True)
    out = []
    for sheet in wb.sheetnames:
        if sheet in ("readme", "Datasets", "CreateCTQI"):
            continue
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))
        # BEA writes a non-breaking space inside industry names ("Wood\xa0Products"),
        # which survives into any key built from the name and compares unequal to the
        # ordinary space nobody can see.
        title = str(rows[0][0] or "").replace("\xa0", " ")
        name = title.split(" - ", 1)[1].strip() if " - " in title else title
        years = None
        for r in rows[:12]:
            if r and str(r[0] or "").strip() == "Asset Codes":
                # THE YEARS ARE TEXT, not numbers -- '1925', not 1925. An isinstance
                # check against int/float finds none of them and returns an empty
                # extraction, which the CSV cache then serves forever.
                years = [_year(c) for c in r[2:]]
                years = [y for y in years if y]
                break
        if years is None:
            continue
        for r in rows:
            code = str(r[0] or "").strip()
            if code in TOTALS:
                vals = list(r[2:2 + len(years)])
                for y, v in zip(years, vals):
                    if isinstance(v, (int, float)):
                        out.append((sheet, name, TOTALS[code], y, float(v)))
    wb.close()
    return out


def table(which):
    """{(industry, asset): {year: value}} plus a code -> name map."""
    csv_path = CACHE / f"beafa_{which}.csv"
    if not csv_path.exists():
        recs = _extract(which)
        # NEVER CACHE AN EMPTY EXTRACTION. This repo has now been bitten twice by a
        # failed parse being written to cache and served silently ever after -- once by
        # an FDIC request that answered 400 with valid JSON, and once here, by the year
        # headers being text. A parse that finds no industries is a bug, not a result.
        if len(recs) < 1000:
            raise RuntimeError(
                f"extracting {which} produced only {len(recs)} rows, which cannot be "
                f"right for 76 industries x 3 assets x ~100 years. Not caching.")
        with csv_path.open("w", newline="") as fh:
            w = csv.writer(fh)
            w.writerow(["code", "name", "asset", "year", "value"])
            w.writerows(recs)
    out, names = {}, {}
    with csv_path.open() as fh:
        for row in csv.DictReader(fh):
            names[row["code"]] = row["name"]
            out.setdefault((row["code"], row["asset"]), {})[int(row["year"])] = \
                float(row["value"])
    return out, names


#: NAICS-ish prefixes. BEA's manufacturing sheets are 31xx, 32xx and 33xx, and the
#: three-digit codes it uses (336M, 336O, 313T, 315A) are motor vehicles, other
#: transport equipment, textiles and apparel -- all manufacturing, so a prefix test
#: on the first two characters is exactly right and no list has to be maintained.
def is_manufacturing(code):
    return code[:2] in ("31", "32", "33")
